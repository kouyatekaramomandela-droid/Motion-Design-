"""FUTURA-Kouroussa : synthèse stratégique présentée au Maire et au Préfet de Kouroussa (octobre 2026).

15 diapositives 16:9 (1920 x 1080 px = 13,333 x 7,5 in), défilement automatique (~6 min 30), animations natives
PowerPoint enchaînées « Après la précédente », transitions Fondu 0,7 s et Morphose (diapositives 4, 8 et 9),
graphiques natifs modifiables animés par catégorie. Aucun chiffre hors du brief.

Coordonnées en pixels de la maquette 1920 x 1080 (1 px = 6350 EMU, 1 pt = 2 px).
Usage : python3 tools/build_deck.py   (depuis le dossier du projet)
"""
import re
import zipfile
from pathlib import Path

from lxml import etree
from PIL import ImageFont
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_MARKER_STYLE, XL_TICK_MARK
from pptx.enum.lang import MSO_LANGUAGE_ID
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "livrables" / "FUTURA-Kouroussa_Synthese_Maire_Prefet.pptx"
LOGOS = ROOT / "assets" / "logos"
ICONS = ROOT / "assets" / "icons"

NAVY, CREAM, GREEN, GOLD, BRICK = "12264A", "FAF8F3", "1E7B67", "D3A04D", "A9553E"
BLUE, OLIVE, VIOLET, GRAY, INK, WHITE = "3A7CA5", "86A05A", "7A5C9E", "8F98A3", "1B2A41", "FFFFFF"
GOLD_TINT, GREEN_TINT, BLUE_TINT, LINE = "F5ECDA", "E5F0EC", "E4EEF4", "E3DED3"

FOOTER = "FUTURA - KOUROUSSA | Octobre 2026"
NOTE = "Projections prospectives du dossier FUTURA-Kouroussa, hypothèse 4 400 USD l'once, non vérifiées de façon indépendante."
PORTE = "Porté par les Députés de Kouroussa, en partenariat avec Visham SARL et l'APEDK"
LOGO_FILES = ["futura-kouroussa", "republique-de-guinee-armoiries", "assemblee-nationale", "simandou-2040", "guinee"]
TOTAL = 15

FONT_R = "/usr/share/fonts/truetype/crosextra/Carlito-Regular.ttf"
FONT_B = "/usr/share/fonts/truetype/crosextra/Carlito-Bold.ttf"
FIT_REPORT = []


def px(v):
    return Emu(int(round(v * 6350)))


def rgb(h):
    return RGBColor.from_string(h)


# ---------------------------------------------------------------- French typography
def fr(s):
    s = s.replace("'", "\u2019")
    s = re.sub(r" ([:;?!%»])", "\u00a0\\1", s)
    s = s.replace("« ", "«\u00a0")
    for _ in range(2):
        s = re.sub(r"(\d) (\d{3})", "\\1\u00a0\\2", s)
    s = re.sub(r"(\d) (M|Md|oz|USD|ans|jours|localités|membres|point)\b", "\\1\u00a0\\2", s)
    s = re.sub(r"\b(M|Md) USD", "\\1\u00a0USD", s)
    s = re.sub(r"\b(Articles?|alinéa|Mois|Décret|Niveau|NIVEAU) ([\dD])", "\\1\u00a0\\2", s)
    s = re.sub(r"(\d) à (\d)", "\\1\u00a0à\u00a0\\2", s)
    return s


# ---------------------------------------------------------------- text
def R(text, size, color=INK, bold=False, italic=False, cs=None):
    return dict(text=text, size=size, color=color, bold=bold, italic=italic, cs=cs)


def P(runs, align="l", after=0, line=1.0, before=0):
    if isinstance(runs, dict):
        runs = [runs]
    return dict(runs=runs, align=align, after=after, line=line, before=before)


_fonts = {}


def _font(bold, size_pt):
    key = (bold, size_pt)
    if key not in _fonts:
        _fonts[key] = ImageFont.truetype(FONT_B if bold else FONT_R, int(round(size_pt * 2)))
    return _fonts[key]


def _lines(para, width):
    """Lines a paragraph needs at this width (greedy wrap on spaces, NBSP kept), measured with Carlito."""
    text = "".join(fr(r["text"]) for r in para["runs"])
    r0 = max(para["runs"], key=lambda r: r["size"])
    f = _font(any(r["bold"] for r in para["runs"]) and r0["bold"], r0["size"])
    extra = (r0["cs"] or 0) * 2
    n, cur = 1, ""
    for word in text.split(" "):
        cand = word if not cur else cur + " " + word
        if f.getlength(cand) + extra * len(cand) <= width or not cur:
            cur = cand
        else:
            n, cur = n + 1, word
    return n, r0["size"]


def fit_check(name, paras, w, h, inset):
    width = w - inset[0] - inset[2]
    total = 0
    for p in paras:
        n, size = _lines(p, width)
        total += n * size * 2 * 1.2 * p["line"] + (p["after"] + p["before"]) * 2
    avail = h - inset[1] - inset[3]
    if total > avail + 1:
        FIT_REPORT.append(f"  {name}: {total:.0f}px de texte pour {avail:.0f}px")


def fill_tf(tf, paras, anchor="t", inset=(0, 0, 0, 0), wrap=True):
    tf.word_wrap = wrap
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = (px(v) for v in inset)
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    for i, pd in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[pd["align"]]
        p.line_spacing = pd["line"]
        p.space_after = Pt(pd["after"])
        p.space_before = Pt(pd["before"])
        for rd in pd["runs"]:
            r = p.add_run()
            r.text = fr(rd["text"])
            f = r.font
            f.name = "Calibri"
            f.size = Pt(rd["size"])
            f.bold = rd["bold"]
            f.italic = rd["italic"]
            f.color.rgb = rgb(rd["color"])
            f.language_id = MSO_LANGUAGE_ID.FRENCH
            if rd["cs"]:
                r._r.get_or_add_rPr().set("spc", str(int(rd["cs"] * 100)))


def text(slide, x, y, w, h, paras, name, anchor="t", inset=(0, 0, 0, 0)):
    tb = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
    tb.name = name
    fill_tf(tb.text_frame, paras, anchor, inset)
    fit_check(name, paras, w, h, inset)
    return tb


# ---------------------------------------------------------------- shapes
def _shadow(shp, alpha=9, blur=26, dist=6):
    sppr = shp._element.spPr
    for old in sppr.findall(qn("a:effectLst")):
        sppr.remove(old)
    eff = etree.SubElement(sppr, qn("a:effectLst"))
    sh = etree.SubElement(eff, qn("a:outerShdw"), blurRad=str(blur * 6350), dist=str(dist * 6350), dir="5400000",
                          algn="t", rotWithShape="0")
    c = etree.SubElement(sh, qn("a:srgbClr"), val=NAVY)
    etree.SubElement(c, qn("a:alpha"), val=str(alpha * 1000))


def _no_shadow(shp):
    sppr = shp._element.spPr
    for old in sppr.findall(qn("a:effectLst")):
        sppr.remove(old)
    etree.SubElement(sppr, qn("a:effectLst"))


def shape(slide, kind, x, y, w, h, name, fill=None, line=None, lw=0.75, radius=None, shadow=False, alpha=None):
    shp = slide.shapes.add_shape(kind, px(x), px(y), px(w), px(h))
    shp.name = name
    if radius is not None and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        shp.adjustments[0] = min(0.5, radius / min(w, h))
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
        if alpha is not None:
            clr = shp._element.spPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
            etree.SubElement(clr, qn("a:alpha"), val=str(int(alpha * 1000)))
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(lw)
    else:
        shp.line.fill.background()
    if shadow:
        _shadow(shp)
    else:
        _no_shadow(shp)
    return shp


def rect(slide, x, y, w, h, name, fill, **kw):
    return shape(slide, MSO_SHAPE.RECTANGLE, x, y, w, h, name, fill=fill, **kw)


def rrect(slide, x, y, w, h, name, fill=WHITE, radius=16, **kw):
    return shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, name, fill=fill, radius=radius, **kw)


def oval(slide, x, y, d, name, fill=None, line=None, lw=1.0, alpha=None):
    return shape(slide, MSO_SHAPE.OVAL, x, y, d, d, name, fill=fill, line=line, lw=lw, alpha=alpha)


def card(slide, x, y, w, h, accent, name, fill=WHITE, border=True, shadow=True):
    """White rounded card with a colour filet on the left (the recurring motif)."""
    c = rrect(slide, x, y, w, h, name, fill=fill, radius=16, line=LINE if border else None, shadow=shadow)
    f = rect(slide, x, y + 16, 8, h - 32, name + " filet", accent)
    return [c, f]


def shape_text(shp, paras, anchor="m", inset=(0, 0, 0, 0)):
    fill_tf(shp.text_frame, paras, anchor, inset)
    w, h = shp.width / 6350, shp.height / 6350
    fit_check(shp.name, paras, w, h, inset)
    return shp


def badge(slide, cx, cy, d, label, fill, name, color=WHITE, size=None):
    o = oval(slide, cx - d / 2, cy - d / 2, d, name, fill=fill)
    shape_text(o, [P(R(label, size or d * 0.22, color, bold=True), align="c")], "m")
    return o


def icon(slide, key, cx, cy, d, fill, name, dark=False):
    """Pictogram in a coloured circle: circle + icon picture (two shapes)."""
    o = oval(slide, cx - d / 2, cy - d / 2, d, name, fill=fill)
    s = d * 0.52
    pic = slide.shapes.add_picture(str(ICONS / f"{key}{'-nuit' if dark else ''}.png"), px(cx - s / 2), px(cy - s / 2), px(s), px(s))
    pic.name = name + " picto"
    return [o, pic]


def connector(slide, x1, y1, x2, y2, name, color=GRAY, w=2.0, arrow=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, px(x1), px(y1), px(x2), px(y2))
    c.name = name
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(w)
    if arrow:
        ln = c.line._get_or_add_ln()
        etree.SubElement(ln, qn("a:tailEnd"), type="triangle", w="med", len="med")
    return c


def picture(slide, path, x, y, h, name):
    pic = slide.shapes.add_picture(str(path), px(x), px(y), height=px(h))
    pic.name = name
    return pic


# ---------------------------------------------------------------- logos
_logo_ratio = {}


def logo_ratio(name):
    if name not in _logo_ratio:
        from PIL import Image
        im = Image.open(LOGOS / f"{name}.png")
        _logo_ratio[name] = im.width / im.height
    return _logo_ratio[name]


def logo_band(slide, cx, y, h, gap, pad_x, pad_y, prefix, shadow=False, line=None):
    widths = [logo_ratio(n) * h for n in LOGO_FILES]
    bw = sum(widths) + gap * (len(widths) - 1) + 2 * pad_x
    bx = cx - bw / 2
    band = rrect(slide, bx, y, bw, h + 2 * pad_y, f"{prefix} bandeau", fill=WHITE, radius=min(24, (h + 2 * pad_y) / 2),
                 line=line, shadow=shadow)
    shapes, x = [band], bx + pad_x
    for n, w in zip(LOGO_FILES, widths):
        shapes.append(picture(slide, LOGOS / f"{n}.png", x, y + pad_y, h, f"{prefix} logo {n}"))
        x += w + gap
    return shapes, bw


# ---------------------------------------------------------------- slide chrome
def chrome(slide, n, rubric, title, accroche=None, note=False, acc_lines=1):
    """Cream content slide: gold trait, green rubric label, title, optional accroche, footer and logo band."""
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(CREAM)
    rect(slide, 96, 70, 64, 6, "Trait or", GOLD)
    text(slide, 96, 80, 1600, 48, [P(R(rubric, 20, GREEN, bold=True, cs=1.5))], "Rubrique", anchor="m")
    text(slide, 96, 122, 1728, 96, [P(R(title, 40, NAVY, bold=True))], "Titre", anchor="m")
    acc = None
    if accroche:
        acc = text(slide, 96, 214, 1728, 56 * acc_lines, [P(R(accroche, 22, INK), line=1.0)], "Accroche")
    if note:
        text(slide, 96, 950, 1728, 36, [P(R(NOTE, 16, INK, italic=True))], "Note prospective", anchor="m")
    text(slide, 96, 1010, 760, 38, [P(R(FOOTER, 14, INK, bold=True, cs=0.5))], "Pied de page", anchor="m")
    band, bw = logo_band(slide, 0, 0, 34, 14, 18, 11, "Pied", line=LINE)
    dx = 1824 - bw - band[0].left / 6350
    for sh in band:
        sh.left = px(sh.left / 6350 + dx)
        sh.top = px(sh.top / 6350 + 1000)
    text(slide, 1824 - bw - 170, 1010, 150, 38, [P(R(f"{n:02d} / {TOTAL}", 14, NAVY, bold=True), align="r")], "Numéro", anchor="m")
    return acc


def measure(s, size, bold=False):
    return _font(bold, size).getlength(fr(s))


# ---------------------------------------------------------------- charts
def _chart_xml_tweaks(chart, layout):
    cs = chart._chartSpace
    plot_area = cs.chart.plotArea
    old = plot_area.find(qn("c:layout"))
    if old is not None:
        plot_area.remove(old)
    lay = etree.Element(qn("c:layout"))
    ml = etree.SubElement(lay, qn("c:manualLayout"))
    etree.SubElement(ml, qn("c:layoutTarget"), val="inner")
    etree.SubElement(ml, qn("c:xMode"), val="edge")
    etree.SubElement(ml, qn("c:yMode"), val="edge")
    for k, v in zip("xywh", layout):
        etree.SubElement(ml, qn(f"c:{k}"), val=f"{v:.4f}")
    plot_area.insert(0, lay)

    def no_fill_sppr(parent, before_tags):
        old = parent.find(qn("c:spPr"))
        if old is not None:
            parent.remove(old)
        sp = etree.Element(qn("c:spPr"))
        etree.SubElement(sp, qn("a:noFill"))
        ln = etree.SubElement(sp, qn("a:ln"))
        etree.SubElement(ln, qn("a:noFill"))
        idx = len(parent)
        for i, child in enumerate(parent):
            if etree.QName(child).localname in before_tags:
                idx = i
                break
        parent.insert(idx, sp)

    no_fill_sppr(cs, {"txPr", "externalData", "printSettings", "userShapes", "extLst"})
    no_fill_sppr(plot_area, {"extLst"})
    rc = cs.find(qn("c:roundedCorners"))
    if rc is None:
        rc = etree.Element(qn("c:roundedCorners"))
        cs.insert(0, rc)
    rc.set("val", "0")


def bar_chart(slide, x, y, w, h, cats, vals, colors, vmax, gap, name):
    cd = CategoryChartData(number_format="0.0")
    cd.categories = [fr(c) for c in cats]
    cd.add_series("Valeurs", vals)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, px(x), px(y), px(w), px(h), cd)
    gf.name = name
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = False
    ch.font.name = "Calibri"
    ch.font.size = Pt(20)
    plot = ch.plots[0]
    plot.gap_width = gap
    plot.overlap = 0
    plot.vary_by_categories = False
    ser = plot.series[0]
    ser.invert_if_negative = False
    for i, c in enumerate(colors):
        pt = ser.points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = rgb(c)
        pt.format.line.fill.background()
    ca, va = ch.category_axis, ch.value_axis
    ca.reverse_order = True
    ca.has_major_gridlines = False
    ca.major_tick_mark = XL_TICK_MARK.NONE
    ca.visible = False
    va.minimum_scale, va.maximum_scale = 0, vmax
    va.has_major_gridlines = False
    va.visible = False
    _chart_xml_tweaks(ch, (0, 0, 1, 1))
    return gf


def line_chart(slide, x, y, w, h, cats, vals, vmax, layout, name):
    cd = CategoryChartData(number_format="0.0")
    cd.categories = cats
    cd.add_series("Contribution annuelle (M USD)", vals)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, px(x), px(y), px(w), px(h), cd)
    gf.name = name
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = False
    ch.font.name = "Calibri"
    ch.font.size = Pt(20)
    ch.font.color.rgb = rgb(INK)
    ser = ch.plots[0].series[0]
    ser.smooth = False
    ser.format.line.color.rgb = rgb(GREEN)
    ser.format.line.width = Pt(4)
    ser.marker.style = XL_MARKER_STYLE.CIRCLE
    ser.marker.size = 14
    ser.marker.format.fill.solid()
    ser.marker.format.fill.fore_color.rgb = rgb(GREEN)
    ser.marker.format.line.color.rgb = rgb(WHITE)
    ser.marker.format.line.width = Pt(2.5)
    ca, va = ch.category_axis, ch.value_axis
    ca.major_tick_mark = XL_TICK_MARK.NONE
    ca.has_major_gridlines = False
    ca.tick_labels.font.size = Pt(20)
    ca.tick_labels.font.bold = True
    ca.tick_labels.font.color.rgb = rgb(INK)
    ca.format.line.color.rgb = rgb("CFC8BA")
    ca.format.line.width = Pt(1.25)
    va.minimum_scale, va.maximum_scale = 0, vmax
    va.has_major_gridlines = False
    va.visible = False
    _chart_xml_tweaks(ch, layout)
    return gf


# ---------------------------------------------------------------- tables
def table(slide, x, y, colw, rows, rowh, name, header_fill=NAVY, zebra=("FFFFFF", "F3EFE6"), aligns=None):
    w = sum(colw)
    h = rowh[0] + rowh[1] * (len(rows) - 1)
    gf = slide.shapes.add_table(len(rows), len(colw), px(x), px(y), px(w), px(h))
    gf.name = name
    tbl = gf.table
    tbl.horz_banding = False
    tbl.first_row = True
    for j, cw in enumerate(colw):
        tbl.columns[j].width = px(cw)
    for i in range(len(rows)):
        tbl.rows[i].height = px(rowh[0] if i == 0 else rowh[1])
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(header_fill if i == 0 else zebra[(i - 1) % 2])
            al = (aligns[j] if aligns else "l")
            bold = i == 0 or j == 0 or j == len(row) - 1
            color = WHITE if i == 0 else (NAVY if j == len(row) - 1 else INK)
            fill_tf(cell.text_frame, [P(R(val, 20, color, bold=bold), align=al)], "m", (16, 0, 16, 0))
            cell.margin_left = cell.margin_right = px(16)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    return gf


# ---------------------------------------------------------------- animation timeline
class Seq:
    """Entrance sequence of one slide: steps start « after previous » or « with previous » (+ delay)."""

    def __init__(self):
        self.steps = []

    def add(self, targets, fx="fade", dur=0.6, start="after", delay=0.0, dirn=None, iterate=None):
        if not isinstance(targets, (list, tuple)) or (len(targets) == 2 and isinstance(targets[1], int)):
            targets = [targets]
        self.steps.append(dict(targets=list(targets), fx=fx, dur=dur, start=start, delay=delay, dirn=dirn, iterate=iterate))
        return self


NS = {
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
    "p14": "http://schemas.microsoft.com/office/powerpoint/2010/main",
    "p159": "http://schemas.microsoft.com/office/powerpoint/2015/09/main",
}
WIPE = {"left": (8, "wipe(right)"), "top": (1, "wipe(down)"), "bottom": (4, "wipe(up)"), "right": (2, "wipe(left)")}


class Ids:
    def __init__(self, start):
        self.n = start

    def __call__(self):
        self.n += 1
        return self.n


def _tgt(t):
    if isinstance(t, tuple):
        gf, cat = t
        return (f'<p:spTgt spid="{gf.shape_id}"><p:graphicEl><a:chart seriesIdx="-4" categoryIdx="{cat}" '
                f'bldStep="category"/></p:graphicEl></p:spTgt>')
    return f'<p:spTgt spid="{t.shape_id}"/>'


def _effect(ids, t, fx, dur, delay, node_type, dirn, iterate):
    tg = _tgt(t)
    d = int(round(dur * 1000))
    preset = {"fade": (10, 0), "wipe": (22, WIPE[dirn or "left"][0]), "float": (42, 0), "zoom": (53, 16)}[fx]
    cid = ids()
    it = f'<p:iterate type="wd"><p:tmAbs val="{iterate}"/></p:iterate>' if iterate else ""
    set_vis = (f'<p:set><p:cBhvr><p:cTn id="{ids()}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst>'
               f'</p:cTn><p:tgtEl>{tg}</p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst>'
               f'</p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>')

    def anim(attr, frm, to, decel=False):
        dc = ' decel="100000"' if decel else ""
        return (f'<p:anim calcmode="lin" valueType="num"><p:cBhvr><p:cTn id="{ids()}" dur="{d}" fill="hold"{dc}/>'
                f'<p:tgtEl>{tg}</p:tgtEl><p:attrNameLst><p:attrName>{attr}</p:attrName></p:attrNameLst></p:cBhvr>'
                f'<p:tavLst><p:tav tm="0"><p:val><p:strVal val="{frm}"/></p:val></p:tav><p:tav tm="100000"><p:val>'
                f'<p:strVal val="{to}"/></p:val></p:tav></p:tavLst></p:anim>')

    def effect(filt):
        return (f'<p:animEffect transition="in" filter="{filt}"><p:cBhvr><p:cTn id="{ids()}" dur="{d}"/>'
                f'<p:tgtEl>{tg}</p:tgtEl></p:cBhvr></p:animEffect>')

    if fx == "fade":
        body = set_vis + effect("fade")
    elif fx == "wipe":
        body = set_vis + effect(WIPE[dirn or "left"][1])
    elif fx == "float":  # entrée par le bas, courte distance, avec fondu (effet natif « Flottant entrant »)
        body = set_vis + effect("fade") + anim("ppt_x", "#ppt_x", "#ppt_x", True) + anim("ppt_y", "#ppt_y+.04", "#ppt_y", True)
    else:  # zoom avant discret (effet natif « Zoom », départ à 85 %)
        body = set_vis + anim("ppt_w", "#ppt_w*0.85", "#ppt_w") + anim("ppt_h", "#ppt_h*0.85", "#ppt_h") + effect("fade")
    grp = "" if isinstance(t, tuple) else ' grpId="0"'
    if isinstance(t, tuple):
        grp = ' grpId="0"'
    return (f'<p:par><p:cTn id="{cid}" presetID="{preset[0]}" presetClass="entr" presetSubtype="{preset[1]}" fill="hold"'
            f'{grp} nodeType="{node_type}"><p:stCondLst><p:cond delay="{int(round(delay * 1000))}"/></p:stCondLst>{it}'
            f'<p:childTnLst>{body}</p:childTnLst></p:cTn></p:par>')


def _words(shp):
    return len(" ".join(p.text for p in shp.text_frame.paragraphs).replace("\u00a0", " ").split())


def timing_xml(seq):
    ids = Ids(3)
    groups, cur_end = [], 0.0
    for st in seq.steps:
        extra = 0.0
        if st["iterate"]:
            extra = (max(_words(t) for t in st["targets"]) - 1) * st["iterate"] / 1000
        if st["start"] == "after" or not groups:
            groups.append(dict(start=cur_end, effects=[]))
        g = groups[-1]
        for k, t in enumerate(st["targets"]):
            first = st is seq.steps[0] or (st["start"] == "after" and k == 0)
            g["effects"].append((t, st, "afterEffect" if (st["start"] == "after" and k == 0) or (first and k == 0) else "withEffect"))
        cur_end = max(cur_end, g["start"] + st["delay"] + st["dur"] + extra)
    pars = []
    for g in groups:
        gid = ids()
        eff = "".join(_effect(ids, t, st["fx"], st["dur"], st["delay"], nt, st["dirn"], st["iterate"]) for t, st, nt in g["effects"])
        pars.append(f'<p:par><p:cTn id="{gid}" fill="hold"><p:stCondLst><p:cond delay="{int(round(g["start"] * 1000))}"/>'
                    f'</p:stCondLst><p:childTnLst>{eff}</p:childTnLst></p:cTn></p:par>')
    # build list
    bld, seen = [], set()
    for st in seq.steps:
        for t in st["targets"]:
            if isinstance(t, tuple):
                gf = t[0]
                if gf.shape_id not in seen:
                    seen.add(gf.shape_id)
                    bld.append(f'<p:bldGraphic spid="{gf.shape_id}" grpId="0"><p:bldSub><a:bldChart bld="category" animBg="0"/>'
                               f'</p:bldSub></p:bldGraphic>')
                continue
            if t.shape_id in seen:
                continue
            seen.add(t.shape_id)
            tag = etree.QName(t._element).localname
            if tag == "graphicFrame":
                bld.append(f'<p:bldGraphic spid="{t.shape_id}" grpId="0"><p:bldAsOne/></p:bldGraphic>')
            elif tag == "sp":
                has_fill = t._element.spPr.find(qn("a:solidFill")) is not None
                bld.append(f'<p:bldP spid="{t.shape_id}" grpId="0"{" animBg=\"1\"" if has_fill else ""}/>')
    xml = ('<p:timing><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
           '<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>'
           '<p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0">'
           f'<p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>{"".join(pars)}</p:childTnLst></p:cTn></p:par>'
           '</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond>'
           '</p:prevCondLst><p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond>'
           '</p:nextCondLst></p:seq></p:childTnLst></p:cTn></p:par></p:tnLst>'
           f'<p:bldLst>{"".join(bld)}</p:bldLst></p:timing>')
    return xml, cur_end


def transition_xml(kind, adv_ms):
    adv = f' advTm="{adv_ms}"' if adv_ms else ""
    if kind == "morph":
        return (f'<mc:AlternateContent xmlns:mc="{NS["mc"]}"><mc:Choice xmlns:p159="{NS["p159"]}" xmlns:p14="{NS["p14"]}" '
                f'Requires="p159"><p:transition spd="slow" p14:dur="1000"{adv}><p159:morph option="byObject"/></p:transition>'
                f'</mc:Choice><mc:Fallback><p:transition spd="slow"{adv}><p:fade/></p:transition></mc:Fallback>'
                f'</mc:AlternateContent>')
    return (f'<mc:AlternateContent xmlns:mc="{NS["mc"]}"><mc:Choice xmlns:p14="{NS["p14"]}" Requires="p14">'
            f'<p:transition spd="med" p14:dur="700"{adv}><p:fade/></p:transition></mc:Choice><mc:Fallback>'
            f'<p:transition spd="med"{adv}><p:fade/></p:transition></mc:Fallback></mc:AlternateContent>')


def _parse(fragment):
    decl = " ".join(f'xmlns:{k}="{v}"' for k, v in NS.items() if k in ("p", "a"))
    return list(etree.fromstring(f"<root {decl}>{fragment}</root>"))


SCHEDULE = []


def finish(slide, seq, seconds, kind="fade", notes=""):
    sld = slide._element
    anchor = sld.find(qn("p:clrMapOvr"))
    idx = list(sld).index(anchor) + 1
    tx = _parse(transition_xml(kind, int(seconds * 1000) if seconds else None))
    tm_xml, end = timing_xml(seq) if seq.steps else ("", 0.0)
    els = tx + (_parse(tm_xml) if tm_xml else [])
    for k, el in enumerate(els):
        sld.insert(idx + k, el)
    SCHEDULE.append((len(SCHEDULE) + 1, seconds, round(end, 2), kind))
    if notes:
        slide.notes_slide.notes_text_frame.text = fr(notes)


# ================================================================ slides
def circles(slide, prefix):
    out = []
    for i, r in enumerate((250, 400, 570)):
        out.append(oval(slide, 1560 - r, 300 - r, 2 * r, f"{prefix} cercle {i + 1}", line=GOLD, lw=1.25, fill=None))
    for c in out:  # transparency on the outline
        ln = c._element.spPr.find(qn("a:ln"))
        clr = ln.find(qn("a:solidFill")).find(qn("a:srgbClr"))
        etree.SubElement(clr, qn("a:alpha"), val="55000")
    return out


def dark(slide):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(NAVY)


def s01(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    dark(s)
    cs = circles(s, "Titre")
    trait = rect(s, 120, 168, 72, 6, "Trait or", GOLD)
    title = text(s, 120, 186, 1600, 160, [P(R("FUTURA KOUROUSSA", 72, WHITE, bold=True, cs=1))], "Titre", anchor="m")
    sub = text(s, 120, 350, 1240, 150, [P(R("Fonds d'Utilisation Transparente des Ressources Aurifères pour l'Avenir de Kouroussa", 30, CREAM), line=1.0)],
               "Sous-titre")
    men = text(s, 120, 506, 1700, 60, [P(R("Synthèse stratégique, présentée au Maire et au Préfet de Kouroussa, octobre 2026", 24, GOLD, bold=True))],
               "Mention", anchor="m")
    band, _ = logo_band(s, 960, 606, 128, 72, 56, 34, "Titre")
    por = text(s, 120, 842, 1680, 60, [P(R(PORTE, 24, CREAM), align="c")], "Porté par", anchor="m")
    q = Seq()
    q.add(cs[0], "wipe", 0.8, dirn="left").add(cs[1], "wipe", 0.8, "with", 0.25, "left").add(cs[2], "wipe", 0.8, "with", 0.5, "left")
    q.add(trait, "wipe", 0.5, "with", 0.3, "left").add(title, "fade", 0.7, "with", 0.3)
    q.add(sub, "fade", 0.6).add(men, "fade", 0.6).add(band, "fade", 0.7).add(por, "fade", 0.6)
    finish(s, q, 12, notes="Titre. Défilement automatique : 12 s.")


def s02(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    acc = chrome(s, 2, "01 / L'ESSENTIEL", "Pourquoi FUTURA-Kouroussa ?",
                 "Transformer une richesse minière temporaire en services publics, activités économiques et perspectives durables pour les habitants.",
                 acc_lines=2)
    cst = card(s, 96, 342, 1728, 236, GOLD, "s02 constat", fill=GOLD_TINT, border=False, shadow=False)
    shape_text(cst[0], [P(R("LE CONSTAT CENTRAL", 20, NAVY, bold=True, cs=1.5), after=6),
                        P(R("Kouroussa accueille trois grands projets aurifères, alors que de nombreux villages manquent d'eau, de soins de santé, "
                            "de routes, d'écoles et d'emplois. Le défi n'est pas seulement de produire de l'or : c'est d'en partager les retombées.", 22, INK), line=1.0)],
               "m", (44, 16, 40, 16))
    q = Seq().add(acc, "fade", 0.6).add(cst, "float", 0.7)
    data = [("1", "CAPTER", "Identifier les contributions existantes, négocier, suivre les versements.", GREEN, "capter"),
            ("2", "GÉRER", "Sécuriser les fonds, tenir des comptes fiables, contrôler et publier chaque dépense.", GOLD, "gerer"),
            ("3", "REDISTRIBUER", "Financer des projets prioritaires choisis avec les communes et les populations.", BLUE, "redistribuer")]
    for i, (num, head, body, col, ico) in enumerate(data):
        x, y, w, h = 96 + i * 588, 604, 552, 356
        c = card(s, x, y, w, h, col, f"s02 carte {head}")
        b = badge(s, x + 70, y + 66, 60, num, col, f"s02 numéro {num}", color=NAVY if col == GOLD else WHITE, size=26)
        hd = text(s, x + 114, y + 42, w - 200, 48, [P(R(head, 26, NAVY, bold=True, cs=1))], f"s02 titre {head}", anchor="m")
        ic = icon(s, ico, x + w - 66, y + 66, 60, GOLD_TINT, f"s02 icône {head}", dark=True)
        bd = text(s, x + 40, y + 124, w - 72, h - 144, [P(R(body, 22, INK), line=1.0)], f"s02 texte {head}")
        q.add(c + [b, hd] + ic + [bd], "float", 0.7, delay=0.15)
    finish(s, q, 25, notes="Pourquoi FUTURA-Kouroussa. Défilement automatique : 25 s.")


def s03(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    acc = chrome(s, 3, "02 / LES PORTEURS", "Qui porte le projet", "Trois partenaires, trois rôles complémentaires.")
    q = Seq().add(acc, "fade", 0.6)
    data = [
        ("Les Députés de Kouroussa", "L'Honorable Lamine KOUYATÉ et l'Honorable Dre Hadja Djoran KEITA",
         "Portage politique et mission parlementaire de septembre 2026 dans les 15 localités.", GOLD, "deputes"),
        ("Visham SARL", None,
         "Partenaire et actionnaire fondateur pendant le montage, avec des règles écrites de prévention des conflits d'intérêts.", GREEN, "visham"),
        ("APEDK", "Association des Personnes Engagées pour le Développement de Kouroussa",
         "Relais auprès des communautés et de la société civile.", BLUE, "apedk"),
    ]
    for i, (head, sub, body, col, ico) in enumerate(data):
        x, y, w, h = 96 + i * 588, 290, 552, 580
        c = card(s, x, y, w, h, col, f"s03 carte {i + 1}")
        ic = icon(s, ico, x + 88, y + 88, 96, col, f"s03 icône {i + 1}", dark=(col == GOLD))
        hd = text(s, x + 152, y + 36, w - 176, 106, [P(R(head, 26, NAVY, bold=True), line=1.0)], f"s03 titre {i + 1}", anchor="m")
        paras = []
        if sub:
            paras.append(P(R(sub, 22, col if col != GOLD else NAVY, bold=True), after=14, line=1.0))
        paras.append(P(R(body, 22, INK), line=1.0))
        tx = text(s, x + 40, y + 170, w - 72, h - 194, paras, f"s03 texte {i + 1}")
        q.add(c + ic + [hd, tx], "float", 0.7, delay=0.15)
    finish(s, q, 20, notes="Qui porte le projet. Défilement automatique : 20 s.")


def s04(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    chrome(s, 4, "03 / LES MINES", "Trois mines, trois profils", note=True)
    rows = [
        ("Kouroussa Gold Mine", "Mansa Resources (société privée enregistrée à Dubaï)",
         "En production depuis novembre 2024", "100 000 à 140 000 oz par an", GOLD, 528, "528 M USD"),
        ("Kiniero", "Predictive Discovery", "En production depuis février 2026", "154 000 à 174 000 oz par an", GREEN, 721.6, "721,6 M USD"),
        ("Bankan", "Predictive Discovery", "Démarrage prévu en 2028", "250 000 à 269 000 oz par an", BRICK, 1142, "1 142 M USD"),
    ]
    y0, hc, g = 240, 170, 12
    bx, bw = 1262, 556
    lab_w = measure("1 142 M USD", 24, True) + 24
    vmax = 1142 * bw / (bw - lab_w - 10)
    cards = [card(s, 96, y0 + i * (hc + g), 1728, hc, r[4], f"s04 carte {i + 1}") for i, r in enumerate(rows)]
    gf = bar_chart(s, bx, y0 - g / 2, bw, 3 * (hc + g), [r[0] for r in rows], [r[5] for r in rows], [r[4] for r in rows],
                   vmax, int(round(((hc + g) / 52 - 1) * 100)), "s04 graphique CA")
    q = Seq()
    for i, (name, op, status, prod, col, val, lab) in enumerate(rows):
        y = y0 + i * (hc + g)
        c = cards[i]
        stc = {GOLD: NAVY, GREEN: GREEN, BRICK: BRICK}[col]
        tx = text(s, 136, y + 8, 1110, hc - 16, [
            P(R(name, 26, NAVY, bold=True)),
            P([R("Opérateur : ", 20, INK, bold=True), R(op, 20, INK)]),
            P([R(status, 20, stc, bold=True), R(" · " + prod, 20, INK)]),
        ], f"s04 texte {i + 1}", anchor="m")
        end = bx + bw * val / vmax
        vl = text(s, end + 12, y + hc / 2 - 26, lab_w + 10, 52, [P(R(lab, 24, NAVY, bold=True))], f"s04 valeur {i + 1}", anchor="m")
        q.add(c + [tx], "float", 0.7, delay=0.1)
        q.add((gf, i), "wipe", 0.7, "with", 0.35, "left")
        q.add(vl, "fade", 0.5, "with", 0.9)
    cap = text(s, bx, y0 + 14, bw, 40, [P(R("CA ANNUEL MOYEN ESTIMÉ", 20, GREEN, bold=True, cs=1.5))], "s04 légende graphique", anchor="m")
    q.steps.insert(1, dict(targets=[cap], fx="fade", dur=0.5, start="with", delay=0.2, dirn=None, iterate=None))
    key = rrect(s, 96, 792, 1728, 146, "s04 chiffre clé", fill=GOLD_TINT, radius=18)
    kf = text(s, 136, 792, 680, 146, [P([R("2,392 Md USD", 44, NAVY, bold=True)], line=0.9),
                                      P(R("de CA annuel moyen estimé", 20, INK, bold=True))],
              "s04 chiffre clé texte", anchor="m")
    hyp = text(s, 830, 792, 970, 146, [P(R("Hypothèse de 4 400 USD l'once (septembre 2026) : estimation, et non chiffre réalisé ou audité.", 20, INK), line=1.0)],
               "s04 hypothèse", anchor="m")
    q.add([key, kf], "zoom", 0.7).add(hyp, "fade", 0.6)
    finish(s, q, 30, "morph", notes="Trois mines, trois profils. Défilement automatique : 30 s. Transition Morphose.")


def s05(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    chrome(s, 5, "04 / LE CADRE LÉGAL", "Ce que prévoit la loi")
    ban = rrect(s, 96, 240, 1728, 196, "s05 bloc 1 %", fill=NAVY, radius=18)
    pct = text(s, 136, 240, 330, 196, [P(R("1 %", 96, GOLD, bold=True))], "s05 1 %", anchor="m")
    btx = text(s, 480, 252, 1310, 172, [
        P(R("du chiffre d'affaires brut hors taxe", 28, WHITE, bold=True), after=4),
        P(R("Contribution au Développement Local pour l'or, substance de Catégorie 2 (Code Minier, Article 130, alinéa 4)", 22, CREAM), line=1.0),
    ], "s05 bloc texte", anchor="m")
    q = Seq().add([ban, pct, btx], "zoom", 0.7)
    rows = [
        ("Article 130", "Convention de Développement Local"),
        ("Décret D/2017/285", "Compte dédié, publication sous cinq jours ouvrés, rapport ITIE avant le 15 avril"),
        ("Article 177", "Contribution déductible : coût réel net de 0,7 % à 1,05 %"),
        ("Articles 107 à 109", "Emploi et fournisseurs locaux, formation"),
        ("Article 160", "Redevance superficiaire"),
        ("Articles 124, 125 et\u00a0165", "Populations déplacées, 15 % de certaines recettes minières aux collectivités"),
    ]
    y = 456
    for i, (lab, desc) in enumerate(rows):
        nl, _ = _lines(P(R(lab, 20, GREEN, bold=True)), 344)
        nd, _ = _lines(P(R(desc, 20, INK)), 700)
        hh = max(nl, nd) * 48 + 12
        ln = connector(s, 96, y + hh, 1150, y + hh, f"s05 filet {i + 1}", color=LINE, w=1.0)
        lt = text(s, 96, y + 8, 344, hh - 10, [P(R(lab, 20, GREEN, bold=True), line=1.0)], f"s05 article {i + 1}")
        dt = text(s, 450, y + 8, 700, hh - 10, [P(R(desc, 20, INK), line=1.0)], f"s05 texte {i + 1}")
        q.add([lt, dt, ln], "float", 0.5, delay=0.05)
        y += hh
    box = card(s, 1176, 456, 648, 506, GOLD, "s05 encadré", fill=GOLD_TINT, border=False, shadow=False)
    shape_text(box[0], [
        P(R("À PLEIN RÉGIME", 20, GREEN, bold=True, cs=1.5), after=6),
        P(R("1 % de 2 392 M USD, soit environ 23,9 M USD par an", 22, NAVY, bold=True), after=8, line=1.0),
        P(R("dont 35 % (environ 8,4 M USD) reviennent aux communes qui abritent les mines.", 20, INK), after=14, line=1.0),
        P(R("FUTURA complète le FODEL national, il ne le concurrence pas.", 20, NAVY, bold=True), line=1.0),
    ], "m", (40, 16, 32, 16))
    q.add(box, "float", 0.7)
    finish(s, q, 30, notes="Ce que prévoit la loi. Défilement automatique : 30 s.")


def s06(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    acc = chrome(s, 6, "05 / LE MÉCANISME", "Un mécanisme à deux niveaux",
                 "Un plancher légal et une majoration proposée selon le résultat.")
    q = Seq().add(acc, "fade", 0.6)

    def level(x, col, tag, val, body, tagc):
        c = card(s, x, 288, 760, 462, col, f"s06 carte {tag}")
        pill = rrect(s, x + 44, 314, 214, 48, f"s06 étiquette {tag}", fill=col, radius=24)
        shape_text(pill, [P(R(tag, 20, tagc, bold=True, cs=1.5), align="c")], "m")
        v = text(s, x + 44, 366, 690, 210, [P(R(val, 64, NAVY, bold=True), line=0.95), P(R("du CA brut hors taxe", 26, INK, bold=True))],
                 f"s06 taux {tag}", anchor="m")
        b = text(s, x + 44, 584, 690, 156, [P(R(body, 22, INK), line=1.0)], f"s06 texte {tag}")
        return c + [pill, v, b]

    n1 = level(96, GREEN, "NIVEAU 1", "1,0 %", "Plancher légal, applicable en toutes circonstances.", WHITE)
    n2 = level(1064, GOLD, "NIVEAU 2", "1,5 %", "Proposition si le résultat net certifié de l'exercice précédent atteint ou dépasse 350 M USD.", NAVY)
    arr = connector(s, 876, 540, 1044, 540, "s06 flèche", color=GOLD, w=4, arrow=True)
    thr = text(s, 836, 434, 248, 92, [P(R("seuil de", 20, INK), align="c"), P(R("350 M USD", 22, NAVY, bold=True), align="c")],
               "s06 seuil", anchor="b")
    q.add(n1, "float", 0.7).add(arr, "wipe", 0.6, dirn="left").add(thr, "fade", 0.5, "with", 0.2).add(n2, "float", 0.7)
    pills = [("Vérification sur les comptes certifiés (Article 187)", 560), ("Copie transmise à FUTURA", 380),
             ("Arbitrage d'un cabinet indépendant sous 60 jours en cas de désaccord", 740)]
    x = 96
    for i, (ptxt, pw) in enumerate(pills):
        pl = rrect(s, x, 774, pw, 124, f"s06 pastille {i + 1}", fill=WHITE, radius=40, line=LINE, shadow=True)
        shape_text(pl, [P(R(ptxt, 20, INK), line=1.0)], "m", (34, 10, 28, 10))
        q.add(pl, "fade", 0.5, "after" if i == 0 else "with", 0.25 * i)
        x += pw + 24
    men = text(s, 96, 912, 1728, 50, [P(R("La majoration de 0,5 point est une proposition à négocier, pas un versement acquis.", 22, BRICK, bold=True, italic=True))],
               "s06 mention", anchor="m")
    q.add(men, "fade", 0.6)
    finish(s, q, 25, notes="Un mécanisme à deux niveaux. Défilement automatique : 25 s.")


def s07(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    acc = chrome(s, 7, "06 / LES SIMULATIONS", "Simulations 2028 et 2029", "Contribution simulée par mine, en millions USD.", note=True)
    q = Seq().add(acc, "fade", 0.6)
    head = ["Mine", "CA", "Résultat net", "Taux", "Contribution"]
    t28 = [head, ["KGM", "572", "165", "1 %", "5,7"], ["Kiniero", "726", "395", "1,5 %", "10,9"], ["Bankan", "880", "710", "1,5 %", "13,2"]]
    t29 = [head, ["KGM", "484", "130", "1 %", "4,8"], ["Kiniero", "678", "360", "1,5 %", "10,2"], ["Bankan", "1 184", "960", "1,5 %", "17,8"]]
    colw = [260, 200, 260, 200, 280]
    al = ["l", "r", "r", "r", "r"]
    bands = []
    for k, (yr, rows, tot) in enumerate([("2028", t28, "29,8"), ("2029", t29, "32,8")]):
        y = 286 + k * 330
        yl = text(s, 96, y, 1200, 50, [P(R(yr, 28, NAVY, bold=True))], f"s07 année {yr}", anchor="m")
        tb = table(s, 96, y + 54, colw, rows, (58, 60), f"s07 tableau {yr}", aligns=al)
        band = rrect(s, 1336, y + 54, 488, 238, f"s07 total {yr}", fill=NAVY, radius=18)
        shape_text(band, [P(R(f"TOTAL {yr}", 24, WHITE, bold=True, cs=1.5), align="c", after=4),
                          P(R(f"{tot} M USD", 48, GOLD, bold=True), align="c")], "m")
        q.add(yl, "fade", 0.5).add(tb, "wipe", 0.8, "with", 0.2, "top")
        bands.append(band)
    q.add(bands[0], "float", 0.6).add(bands[1], "float", 0.6, delay=0.1)
    finish(s, q, 25, notes="Simulations 2028 et 2029. Défilement automatique : 25 s.")


def s08(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    acc = chrome(s, 8, "07 / L'HORIZON FINANCIER", "Horizon financier",
                 "Contributions annuelles projetées : les projections commencent en 2028, aucun cumul avant.", note=True)
    q = Seq().add(acc, "fade", 0.6)
    cats = ["2028", "2029", "2030", "2031", "2032"]
    vals = [29.8, 32.8, 32.2, 30.6, 22.7]
    X, Y, W, H = 96, 288, 1728, 252
    lay = (0.03, 0.40, 0.94, 0.40)
    vmax = 40
    gf = line_chart(s, X, Y, W, H, cats, vals, vmax, lay, "s08 courbe")
    band = W * lay[2] / 5
    pills = {0: "2028", 1: "2029", 4: "2032"}
    for i, v in enumerate(vals):
        cx = X + W * lay[0] + band * (i + 0.5)
        cy = Y + H * (lay[1] + lay[3] * (1 - v / vmax))
        vtxt = f"{v:.1f}".replace(".", ",")
        if i in pills:
            p = rrect(s, cx - 140, cy - 122, 280, 104, f"s08 pastille {cats[i]}", fill=NAVY, radius=24)
            shape_text(p, [P(R(f"{vtxt} M USD", 24, WHITE, bold=True), align="c"), P(R(f"en {pills[i]}", 20, CREAM), align="c")], "m")
            lab = p
        else:
            lab = text(s, cx - 80, cy - 62, 160, 44, [P(R(vtxt, 22, GREEN, bold=True), align="c")], f"s08 valeur {cats[i]}", anchor="m")
        q.add((gf, i), "wipe", 0.6, dirn="left").add(lab, "fade", 0.5, "with", 0.3)
    box = card(s, 96, 560, 840, 372, GOLD, "s08 encadré", fill=GOLD_TINT, border=False, shadow=False)
    shape_text(box[0], [
        P(R("148,1 M USD", 48, NAVY, bold=True)),
        P(R("cumulés sur cinq exercices, 2028 à 2032", 22, INK, bold=True), after=14),
        P(R("Sur la durée de vie des mines : environ 178 à 265 M USD.", 22, INK), line=1.0),
    ], "m", (44, 20, 36, 20))
    q.add(box, "float", 0.7)
    cards = [("Ce qui fait varier le montant", "Cours de l'or, volumes, résultat net certifié, négociation.", BRICK),
             ("Ce qui doit être validé", "Assiette disponible, clé de répartition, conventions signées, calendrier de versement.", GREEN)]
    for i, (hd, body, col) in enumerate(cards):
        y = 560 + i * 192
        c = card(s, 960, y, 864, 180, col, f"s08 carte {i + 1}")
        tx = text(s, 1000, y + 10, 800, 160, [P(R(hd, 22, NAVY, bold=True), after=4), P(R(body, 20, INK), line=1.0)],
                  f"s08 carte texte {i + 1}", anchor="m")
        q.add(c + [tx], "float", 0.7, delay=0.1)
    finish(s, q, 30, "morph", notes="Horizon financier. Défilement automatique : 30 s. Transition Morphose.")


def s09(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    acc = chrome(s, 9, "08 / LES PRIORITÉS", "À quoi serviraient les ressources ?",
                 "Répartition proposée des ressources : huit postes, 100 % du total.")
    q = Seq().add(acc, "fade", 0.6)
    rows = [("Infrastructures, eau, électricité, routes", 25, GOLD), ("Santé et action sociale", 20, GREEN),
            ("Éducation et formation", 15, BLUE), ("Agriculture et sécurité alimentaire", 10, OLIVE),
            ("Emploi et développement économique", 10, BRICK), ("Fonds Après-Mine", 10, VIOLET),
            ("Obligations d'État, épargne des générations futures", 7, NAVY), ("Fonctionnement et réserve d'urgence", 3, GRAY)]
    y0, rh, bar = 290, 80, 24
    X, W, vmax = 96, 1000, 27
    gf = bar_chart(s, X, y0 + 60 - rh / 2, W, rh * len(rows), [r[0] for r in rows], [r[1] for r in rows], [r[2] for r in rows],
                   vmax, int(round((rh / bar - 1) * 100)), "s09 graphique répartition")
    for i, (lab, v, col) in enumerate(rows):
        y = y0 + i * rh
        lt = text(s, X, y, 1060, 44, [P(R(lab, 20, INK, bold=True))], f"s09 poste {i + 1}", anchor="b")
        end = X + W * v / vmax
        vt = text(s, end + 12, y + 60 - 22, 110, 44, [P(R(f"{v} %", 22, NAVY, bold=True))], f"s09 part {i + 1}", anchor="m")
        q.add(lt, "fade", 0.4, delay=0.05).add((gf, i), "wipe", 0.6, "with", 0.1, "left").add(vt, "fade", 0.4, "with", 0.5)
    hd = text(s, 1236, 286, 588, 92, [P(R("HUIT AXES ISSUS DE LA MISSION PARLEMENTAIRE", 20, GREEN, bold=True, cs=1.5), line=1.0)],
              "s09 axes titre", anchor="m")
    q.add(hd, "fade", 0.5)
    axes = [("Routes", "routes", GOLD), ("Eau", "eau", BLUE), ("Santé", "sante", GREEN), ("Éducation", "education", BLUE),
            ("Agriculture", "agriculture", OLIVE), ("Emploi", "emploi", BRICK), ("Femmes", "femmes", VIOLET),
            ("Électricité", "electricite", GOLD)]
    for i, (lab, ico, col) in enumerate(axes):
        cx, cy = 1236 + (i % 2) * 300, 392 + (i // 2) * 136
        pl = rrect(s, cx, cy, 288, 112, f"s09 axe {i + 1}", fill=WHITE, radius=24, line=LINE, shadow=True)
        ic = icon(s, ico, cx + 48, cy + 56, 56, col, f"s09 axe icône {i + 1}", dark=(col in (GOLD, OLIVE)))
        tx = text(s, cx + 86, cy, 198, 112, [P(R(lab, 20, NAVY, bold=True))], f"s09 axe texte {i + 1}", anchor="m")
        q.add([pl] + ic + [tx], "fade", 0.5, "with" if i else "after", 0.12 * i)
    finish(s, q, 30, "morph", notes="À quoi serviraient les ressources. Défilement automatique : 30 s. Transition Morphose.")


def s10(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    acc = chrome(s, 10, "09 / LA GOUVERNANCE", "Qui décide, qui exécute, qui contrôle", "Cinq organes, des rôles séparés.")
    q = Seq().add(acc, "fade", 0.6)
    tagc = {GOLD: NAVY}

    def organ(x, y, w, h, col, tag, title, body, name):
        c = card(s, x, y, w, h, col, f"s10 {name}")
        paras = [P([R(tag, 20, tagc.get(col, col), bold=True, cs=1.5), R("   " + title, 24, NAVY, bold=True)], line=1.0, after=4)]
        if body:
            paras.append(P(R(body, 20, INK), line=1.0))
        tx = text(s, x + 36, y + 10, w - 64, h - 20, paras, f"s10 texte {name}", anchor="m")
        return c + [tx]

    copis = organ(360, 288, 1200, 118, GOLD, "ORIENTE", "Comité de Pilotage Stratégique (CoPiS)", "Réunion annuelle", "CoPiS")
    c1 = connector(s, 960, 406, 960, 426, "s10 lien 1", color=GRAY, w=2)
    ca = card(s, 96, 426, 1728, 178, GREEN, "s10 CA")
    catx = text(s, 132, 436, 1660, 54, [P([R("DÉCIDE", 20, GREEN, bold=True, cs=1.5), R("   Conseil d'Administration", 24, NAVY, bold=True),
                                          R(" · 11 membres, quatre collèges", 22, INK)])], "s10 texte CA", anchor="m")
    chips = [("3", "Collectivités locales", GREEN), ("3", "Sociétés minières", GOLD), ("3", "Administrateurs indépendants", BLUE), ("2", "État", VIOLET)]
    chip_shapes = []
    for i, (n, lab, col) in enumerate(chips):
        cx = 132 + i * 416
        ch = rrect(s, cx, 498, 396, 88, f"s10 collège {i + 1}", fill=CREAM, radius=14)
        nb = text(s, cx + 14, 498, 64, 88, [P(R(n, 40, col if col != GOLD else NAVY, bold=True), align="c")], f"s10 collège nombre {i + 1}", anchor="m")
        lb = text(s, cx + 84, 498, 300, 88, [P(R(lab, 20, INK, bold=True), line=1.0)], f"s10 collège texte {i + 1}", anchor="m")
        chip_shapes.append([ch, nb, lb])
    c2 = connector(s, 518, 604, 518, 624, "s10 lien 2", color=GRAY, w=2)
    c3 = connector(s, 1402, 604, 1402, 624, "s10 lien 3", color=GRAY, w=2)
    dg = organ(96, 624, 844, 172, BLUE, "EXÉCUTE", "Direction Générale", "Le Directeur Général n'est pas administrateur.", "DG")
    au = organ(980, 624, 844, 172, BRICK, "CONTRÔLE", "Comité d'Audit et de Surveillance", "Rapport public avant le 30 juin.", "Audit")
    c4 = connector(s, 960, 796, 960, 816, "s10 lien 4", color=GRAY, w=2)
    asm = organ(96, 816, 1728, 140, VIOLET, "EXPRIME LES BESOINS", "Assemblée Communautaire",
                "Valide les priorités, élit le Comité d'Audit ; au moins un tiers de femmes.", "Assemblée")
    q.add(copis, "float", 0.6).add(c1, "wipe", 0.5, dirn="top").add(ca + [catx], "float", 0.6)
    for i, cs_ in enumerate(chip_shapes):
        q.add(cs_, "fade", 0.5, "with", 0.3 + 0.15 * i)
    q.add([c2, c3], "wipe", 0.5, dirn="top").add(dg, "float", 0.6).add(au, "float", 0.6, "with", 0.15)
    q.add(c4, "wipe", 0.5, dirn="top").add(asm, "float", 0.6)
    finish(s, q, 30, notes="Qui décide, qui exécute, qui contrôle. Défilement automatique : 30 s.")


def s11(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    acc = chrome(s, 11, "10 / L'INTÉGRITÉ", "Les sept verrous anti-corruption",
                 "Des garanties cumulatives, de la comptabilité à la séparation des pouvoirs.")
    q = Seq().add(acc, "fade", 0.6)
    panel = rrect(s, 1300, 288, 524, 652, "s11 panneau", fill=NAVY, radius=20)
    sh = icon(s, "bouclier", 1562, 432, 176, GOLD, "s11 bouclier", dark=True)
    big = text(s, 1300, 546, 524, 370, [P(R("7", 120, GOLD, bold=True), align="c", line=1.0), P(R("verrous", 30, WHITE, bold=True), align="c")],
               "s11 sept", anchor="t")
    q.add([panel] + sh + [big], "fade", 0.6)
    items = ["Comptes SYSCOHADA certifiés", "Audit externe annuel, par appel d'offres", "Plateforme publique des versements et des projets",
             "Publication des comptes (Journal Officiel, RCCM)", "Canal de signalement accessible à tous",
             "Rotation des mandats : 3 ans, renouvelables une fois", "Séparation stricte des pouvoirs"]
    for i, it in enumerate(items):
        y = 288 + i * 93
        col = GREEN if i % 2 == 0 else GOLD
        b = badge(s, 96 + 30, y + 42, 56, str(i + 1), col, f"s11 numéro {i + 1}", color=WHITE if col == GREEN else NAVY, size=24)
        tx = text(s, 186, y, 1070, 84, [P(R(it, 24, INK, bold=True))], f"s11 verrou {i + 1}", anchor="m")
        ln = connector(s, 96, y + 88, 1256, y + 88, f"s11 filet {i + 1}", color=LINE, w=1.0)
        q.add([b, tx, ln], "float", 0.5, delay=0.05)
    finish(s, q, 25, notes="Les sept verrous anti-corruption. Défilement automatique : 25 s.")


def s12(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    acc = chrome(s, 12, "11 / TRANSPARENCE SUR NOTRE PROPRE POSITION", "Prévention des conflits d'intérêts",
                 "Visham SARL, partenaire et actionnaire fondateur pendant le montage :")
    q = Seq().add(acc, "fade", 0.6)
    rules = ["Aucun siège réservé au Conseil d'Administration", "Déclaration et récusation de toute personne intéressée",
             "Aucun contrat de gré à gré : appel d'offres public et revue du Comité d'Audit",
             "Rôle transitoire jusqu'à l'arrivée de la Direction Générale", "Registre public des intérêts"]
    for i, r_ in enumerate(rules):
        y = 288 + i * 130
        c = card(s, 96, y, 1040, 114, GREEN, f"s12 règle {i + 1}")
        ic = icon(s, "check", 96 + 64, y + 57, 52, GREEN, f"s12 coche {i + 1}")
        tx = text(s, 216, y + 6, 900, 102, [P(R(r_, 22, INK, bold=True), line=1.0)], f"s12 texte {i + 1}", anchor="m")
        q.add(c + ic + [tx], "float", 0.5, delay=0.05)
    box = card(s, 1176, 288, 648, 634, GOLD, "s12 encadré", fill=GOLD_TINT, border=False, shadow=False)
    shape_text(box[0], [
        P(R("Ne pas confondre COPIL et CoPiS", 26, NAVY, bold=True), after=20, line=1.0),
        P([R("Le COPIL", 22, GREEN, bold=True), R(" est le comité de pilotage provisoire de mise en place, jusqu'à l'assemblée générale constitutive.", 22, INK)],
          after=18, line=1.0),
        P([R("Le CoPiS", 22, GREEN, bold=True), R(" est l'organe stratégique permanent de la société constituée.", 22, INK)], line=1.0),
    ], "m", (48, 28, 40, 28))
    q.add(box, "float", 0.7)
    finish(s, q, 25, notes="Transparence sur notre propre position. Défilement automatique : 25 s.")


def s13(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    acc = chrome(s, 13, "12 / LA FEUILLE DE ROUTE", "Feuille de route",
                 "Quatre étapes, de la délibération communale aux premiers projets livrés.")
    q = Seq().add(acc, "fade", 0.6)
    line = rect(s, 168, 374, 1584, 4, "s13 frise", GOLD)
    q.add(line, "wipe", 0.8, dirn="left")
    steps = [("Mois 1 à 3", "Faire front commun", "Délibération communale, concertation, mise en place du COPIL.", GREEN),
             ("Mois 3 à 6", "Sécuriser et négocier", "Dialogue avec les trois sociétés, statuts OHADA, conventions.", GOLD),
             ("Mois 6 à 12", "Créer et contractualiser", "Assemblée constitutive, nominations, compte dédié, signature des conventions.", BLUE),
             ("Mois 12 à 24", "Livrer les premiers projets", "Point d'eau, centre de santé, route critique, premiers résultats publiés.", BRICK)]
    for i, (when, head, items, col) in enumerate(steps):
        x, w = 96 + i * 438, 414
        cx = x + w / 2
        lab = text(s, x, 292, w, 48, [P(R(when.upper(), 20, GREEN if col != GOLD else NAVY, bold=True, cs=1.5), align="c")], f"s13 période {i + 1}", anchor="m")
        node = badge(s, cx, 376, 52, str(i + 1), col, f"s13 étape {i + 1}", color=WHITE if col != GOLD else NAVY, size=22)
        c = card(s, x, 428, w, 532, col, f"s13 carte {i + 1}")
        tx = text(s, x + 36, 452, w - 64, 488, [P(R(head, 26, NAVY, bold=True), after=16, line=1.0), P(R(items, 22, INK), line=1.05)],
                  f"s13 texte {i + 1}")
        q.add([node, lab], "zoom", 0.5).add(c + [tx], "float", 0.6, "with", 0.2)
    finish(s, q, 25, notes="Feuille de route. Défilement automatique : 25 s.")


def s14(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    acc = chrome(s, 14, "13 / NOTRE PROPOSITION", "Ce que nous proposons au Maire et au Préfet",
                 "Le cadre légal existe, le projet l'applique : une décision locale suffit pour démarrer.")
    q = Seq().add(acc, "fade", 0.6)
    cols = [("Mairie de Kouroussa", "Nous proposons de :", GREEN, "mairie", [
        "Voter en session plénière du Conseil Communal la délibération demandant la création de FUTURA-Kouroussa et mandatant le Maire pour négocier.",
        "Constituer le dossier de réclamation auprès de KGM pour les montants de CDL dus depuis novembre 2024.",
        "Désigner les représentants de la commune dans les futurs organes."]),
        ("Préfecture de Kouroussa", "Nous proposons :", BLUE, "prefecture", [
            "Un appui à la démarche.",
            "La présence du Préfet, ou de son représentant, au Comité de Pilotage Stratégique.",
            "Sa présence à la signature solennelle des Conventions de Développement Local."])]
    for k, (head, lead, col, ico, items) in enumerate(cols):
        x, w = 96 + k * 888, 840
        c = card(s, x, 288, w, 672, col, f"s14 colonne {k + 1}")
        ic = icon(s, ico, x + 84, 360, 84, col, f"s14 icône {k + 1}")
        hd = text(s, x + 148, 306, w - 180, 110, [P(R(head, 28, NAVY, bold=True), line=1.0, after=2), P(R(lead, 22, col, bold=True))],
                  f"s14 en-tête {k + 1}", anchor="m")
        q.add(c + ic + [hd], "float", 0.7)
        y = 440
        for i, it in enumerate(items):
            n, _ = _lines(P(R(it, 20, INK)), w - 148)
            hh = n * 48 + 8
            b = badge(s, x + 66, y + 24, 42, str(i + 1), col, f"s14 point {k + 1}.{i + 1}", size=20)
            tx = text(s, x + 108, y, w - 140, hh, [P(R(it, 20, INK), line=1.0)], f"s14 texte {k + 1}.{i + 1}")
            q.add([b, tx], "fade", 0.5, delay=0.05)
            y += hh + 20
    finish(s, q, 30, notes="Ce que nous proposons au Maire et au Préfet. Défilement automatique : 30 s.")


def s15(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    dark(s)
    circles(s, "Conclusion")
    st = text(s, 160, 92, 1600, 452, [
        P(R("Le Code Minier est clair.", 42, WHITE, bold=True), align="c", after=6),
        P(R("Les mines sont en activité.", 42, WHITE, bold=True), align="c", after=6),
        P(R("Les besoins sont documentés.", 42, WHITE, bold=True), align="c", after=6),
        P(R("Il ne manque qu'une décision.", 42, WHITE, bold=True), align="c"),
    ], "Conclusion phrase", anchor="m")
    ready = text(s, 160, 552, 1600, 136, [P(R("FUTURA-Kouroussa est prêt.", 56, GOLD, bold=True), align="c")], "Conclusion prêt", anchor="m")
    band, _ = logo_band(s, 960, 706, 116, 64, 50, 30, "Conclusion")
    por = text(s, 160, 902, 1600, 60, [P(R(PORTE, 24, CREAM), align="c")], "Porté par", anchor="m")
    q = Seq().add(st, "fade", 0.5, iterate=280).add(ready, "zoom", 0.8, delay=0.3).add(band, "fade", 0.7).add(por, "fade", 0.6)
    finish(s, q, None, notes="Conclusion : environ 9 s d'animation puis image fixe (pas d'avancement automatique après la dernière diapositive).")


# ================================================================ theme and packaging
THEME_COLORS = dict(dk1=INK, lt1=WHITE, dk2=NAVY, lt2=CREAM, accent1=GREEN, accent2=GOLD, accent3=BRICK, accent4=BLUE,
                    accent5=OLIVE, accent6=VIOLET, hlink=BLUE, folHlink=VIOLET)


def patch_theme(path):
    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if re.match(r"ppt/theme/theme\d+\.xml", item.filename):
                xml = data.decode("utf8")
                xml = re.sub(r'<a:clrScheme name="[^"]*">', '<a:clrScheme name="FUTURA-Kouroussa">', xml)
                for k, v in THEME_COLORS.items():
                    xml = re.sub(rf"<a:{k}>.*?</a:{k}>", f'<a:{k}><a:srgbClr val="{v}"/></a:{k}>', xml, flags=re.S)
                xml = re.sub(r'(<a:majorFont>\s*<a:latin typeface=")[^"]*"', r'\1Calibri"', xml)
                xml = re.sub(r'(<a:minorFont>\s*<a:latin typeface=")[^"]*"', r'\1Calibri"', xml)
                data = xml.encode("utf8")
            zout.writestr(item, data)
    tmp.replace(path)


def check_text(prs):
    bad = []
    for i, sl in enumerate(prs.slides, 1):
        for shp in sl.shapes:
            texts = []
            if shp.has_text_frame:
                texts.append(shp.text_frame.text)
            if getattr(shp, "has_table", False) and shp.has_table:
                texts += [c.text for r in shp.table.rows for c in r.cells]
            for t in texts:
                if "\u2014" in t or "\u2013" in t:
                    bad.append((i, shp.name, t))
    return bad


def main():
    prs = Presentation()
    prs.slide_width, prs.slide_height = px(1920), px(1080)
    prs.core_properties.title = "FUTURA-Kouroussa, synthèse stratégique"
    prs.core_properties.subject = "Présentation au Maire et au Préfet de Kouroussa, octobre 2026"
    prs.core_properties.author = "FUTURA-Kouroussa"
    prs.core_properties.language = "fr-FR"
    for fn in (s01, s02, s03, s04, s05, s06, s07, s08, s09, s10, s11, s12, s13, s14, s15):
        fn(prs)
    bad = check_text(prs)
    assert not bad, f"tiret long trouvé : {bad}"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    patch_theme(OUT)
    total = sum(sec or 15 for _, sec, _, _ in SCHEDULE) + 0.7 * (TOTAL - 1)
    print(f"écrit {OUT.relative_to(ROOT)} ({OUT.stat().st_size / 1e6:.1f} Mo)")
    print("diapo  durée  fin anim.  transition")
    for n, sec, end, kind in SCHEDULE:
        print(f"  {n:2d}   {sec or 'fixe':>4}   {end:5.1f} s    {kind}")
    print(f"durée totale en défilement automatique : {total:.0f} s ({int(total // 60)} min {int(total % 60):02d})")
    if FIT_REPORT:
        print("textes à vérifier (estimation Carlito) :")
        print("\n".join(FIT_REPORT))


if __name__ == "__main__":
    main()
